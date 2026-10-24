import os
import re
import io
import gc
import math
import time
import hashlib
from datetime import datetime, date, timedelta
from pathlib import Path

import streamlit as st
import pandas as pd
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from google import genai
from google.genai import types

try:
    import streamlit.components.v1 as components
except Exception:
    components = None

try:
    from openpyxl.styles import Font, PatternFill, Alignment
    ESTILOS_EXCEL_OK = True
except Exception:
    ESTILOS_EXCEL_OK = False

# =====================================================================
# CONFIGURACIÓN DE PÁGINA Y ESCUDO ANTI-TRADUCTOR (EVITA ERROR removeChild)
# =====================================================================
st.set_page_config(
    page_title="Bufete Litigante IA Colombia | Auditoría Turbo y Estrategia",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown('<meta name="google" content="notranslate">', unsafe_allow_html=True)
if components is not None:
    try:
        components.html(
            """
            <script>
                try {
                    const doc = window.parent.document;
                    doc.documentElement.setAttribute('lang', 'es');
                    doc.documentElement.setAttribute('translate', 'no');
                    doc.documentElement.classList.add('notranslate');
                    if (doc.body) {
                        doc.body.setAttribute('translate', 'no');
                        doc.body.classList.add('notranslate');
                    }
                    if (!doc.querySelector('meta[name="google"]')) {
                        const m = doc.createElement('meta');
                        m.name = 'google';
                        m.content = 'notranslate';
                        doc.head.appendChild(m);
                    }
                } catch (e) {}
            </script>
            """,
            height=0,
            width=0
        )
    except Exception:
        pass

st.markdown("""
<style>
    .main-header {
        background: linear-gradient(135deg, #0f172a 0%, #1e3a8a 65%, #1d4ed8 100%);
        padding: 1.5rem 2rem;
        border-radius: 12px;
        color: white;
        margin-bottom: 0.9rem;
        border-bottom: 4px solid #f59e0b;
        box-shadow: 0 4px 14px rgba(15, 23, 42, 0.14);
    }
    .status-bar {
        display: flex;
        flex-wrap: wrap;
        gap: 0.65rem;
        margin-bottom: 1.1rem;
    }
    .status-pill {
        background-color: #f1f5f9;
        border: 1px solid #cbd5e1;
        color: #1e293b;
        padding: 0.32rem 0.85rem;
        border-radius: 999px;
        font-size: 0.83rem;
        font-weight: 600;
    }
    .status-pill-ok {
        background-color: #ecfdf5;
        border: 1px solid #6ee7b7;
        color: #065f46;
    }
    .status-pill-warn {
        background-color: #fffbeb;
        border: 1px solid #fcd34d;
        color: #92400e;
    }
    .tool-box {
        background: linear-gradient(180deg, #f8fafc 0%, #f1f5f9 100%);
        border: 1px solid #cbd5e1;
        border-left: 5px solid #1e3a8a;
        padding: 1.15rem;
        border-radius: 10px;
        margin-bottom: 1.1rem;
    }
    .gold-box {
        background-color: #fffbeb;
        border: 1px solid #fde68a;
        border-left: 5px solid #d97706;
        padding: 1.05rem;
        border-radius: 10px;
        color: #78350f;
        margin-bottom: 1rem;
    }
    div[data-testid="stMetric"] {
        background-color: #f8fafc;
        border: 1px solid #e2e8f0;
        border-top: 3px solid #1e3a8a;
        padding: 0.75rem 1rem;
        border-radius: 10px;
    }
</style>
""", unsafe_allow_html=True)

# =====================================================================
# ESTADO DE SESIÓN EN MEMORIA RAM
# =====================================================================
def inicializar_sesion():
    defaults = {
        "expediente_docs": [],
        "cache_hashes": {},
        "actuaciones_df": pd.DataFrame(),
        "catalogo_df": pd.DataFrame(),
        "informe_md": "",
        "estrategia_md": "",
        "blindaje_md": "",
        "historial_chat": [],
        "modelos_cache": {},
        "ultimo_tiempo_lectura": 0.0
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v

inicializar_sesion()

def obtener_api_key_inicial():
    try:
        if "GOOGLE_API_KEY" in st.secrets:
            return str(st.secrets["GOOGLE_API_KEY"]).strip()
    except Exception:
        pass
    return os.environ.get("GOOGLE_API_KEY", "").strip()

def df_a_markdown_seguro(df: pd.DataFrame) -> str:
    if df is None or df.empty:
        return "Sin registros detectados."
    cols = [str(c) for c in df.columns]
    header = "| " + " | ".join(cols) + " |"
    sep = "| " + " | ".join(["---"] * len(cols)) + " |"
    filas = [header, sep]
    for _, row in df.iterrows():
        celdas = [re.sub(r'\s+', ' ', str(row[c]).replace("|", "-")).strip() for c in df.columns]
        filas.append("| " + " | ".join(celdas) + " |")
    return "\n".join(filas)

# =====================================================================
# EXPRESIONES REGULARES PARA METADATOS Y FECHAS
# =====================================================================
RE_FECHA = re.compile(r'\b\d{1,2}[\/\-]\d{1,2}[\/\-](?:\d{4}|\d{2})\b')
RE_RADICADO = re.compile(
    r'(Radicado|Rad\.|Proceso|Expediente|Caso|No\.?)[ \t]*[:#\-]?[ \t]*([0-9]{4,23}(?:[-][0-9]{1,8})*)',
    re.I
)
RE_ACTOR = re.compile(
    r'(Demandante|Actor|Solicitante|Promovente|Accionante|Quejoso|Denunciante|Peticionario|Incidentante)[ \t]*[:\-][ \t]*([^\n\r]{3,120})',
    re.I
)
RE_DEMANDADO = re.compile(
    r'(Demandado|Demandada|Convocado|Entidad|Accionado|Accionada|Investigado|Denunciado|Querellado)[ \t]*[:\-][ \t]*([^\n\r]{3,120})',
    re.I
)
RE_JUEZ = re.compile(
    r'(Juez|Jueza|Magistrad[oa]|Despacho|Juzgado|Superintendencia|Procuradur[ií]a|Personer[ií]a|Fiscal[ií]a|Alcald[ií]a|Inspecci[oó]n|Secretar[ií]a)[ \t]*[:\-]?[ \t]*([^\n\r]{3,120})',
    re.I
)
RE_NUMERO_OFICIO = re.compile(
    r'(OFICIO|Resoluci[oó]n|Auto|Acta)[ \t]+No\.?[ \t]*([0-9A-Za-z\-]{1,20})',
    re.I
)

TIPOS_ACT = [
    'DEMANDA', 'TUTELA', 'DESACATO', 'QUEJA', 'DERECHO DE PETICION', 'CONTESTACION',
    'REPLICA', 'AUTO', 'PROVIDENCIA', 'RESOLUCION', 'SENTENCIA', 'APELACION',
    'RECURSO', 'REPOSICION', 'NOTIFICACION', 'ACTA', 'OFICIO', 'INFORME',
    'PERITAJE', 'LIQUIDACION', 'AUDIENCIA', 'MEMORIAL', 'PRUEBA', 'CONTRATO',
    'PROMESA', 'ESCRITURA', 'LICENCIA', 'QUERELLA', 'DENUNCIA'
]

# =====================================================================
# CALCULADORA OFICIAL DE TÉRMINOS Y FESTIVOS EN COLOMBIA (LEY EMILIANI)
# =====================================================================
def calcular_pascua(year: int) -> date:
    a = year % 19
    b = year // 100
    c = year % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = ((h + l - 7 * m + 114) % 31) + 1
    return date(year, month, day)

def trasladar_a_lunes(d: date) -> date:
    if d.weekday() == 0:
        return d
    return d + timedelta(days=(7 - d.weekday()))

def obtener_festivos_colombia(year: int) -> set:
    fijos = [
        date(year, 1, 1),
        date(year, 5, 1),
        date(year, 7, 20),
        date(year, 8, 7),
        date(year, 12, 8),
        date(year, 12, 25),
    ]
    emiliani = [
        trasladar_a_lunes(date(year, 1, 6)),
        trasladar_a_lunes(date(year, 3, 19)),
        trasladar_a_lunes(date(year, 6, 29)),
        trasladar_a_lunes(date(year, 8, 15)),
        trasladar_a_lunes(date(year, 10, 12)),
        trasladar_a_lunes(date(year, 11, 1)),
        trasladar_a_lunes(date(year, 11, 11)),
    ]
    dom_pascua = calcular_pascua(year)
    religiosos = [
        dom_pascua - timedelta(days=3),
        dom_pascua - timedelta(days=2),
        trasladar_a_lunes(dom_pascua + timedelta(days=39)),
        trasladar_a_lunes(dom_pascua + timedelta(days=60)),
        trasladar_a_lunes(dom_pascua + timedelta(days=68)),
    ]
    return set(fijos + emiliani + religiosos)

def sumar_dias_habiles_colombia(fecha_inicio: date, dias_habiles: int):
    actual = fecha_inicio
    contados = 0
    festivos_omitidos = []
    fines_semana_omitidos = 0
    ruta_dias = []

    while contados < dias_habiles:
        actual += timedelta(days=1)
        festivos_anio = obtener_festivos_colombia(actual.year)
        if actual.weekday() >= 5:
            fines_semana_omitidos += 1
            continue
        if actual in festivos_anio:
            festivos_omitidos.append(actual.strftime("%d/%m/%Y"))
            continue
        contados += 1
        ruta_dias.append({"Día Hábil N°": contados, "Fecha": actual.strftime("%d/%m/%Y")})

    return actual, fines_semana_omitidos, festivos_omitidos, pd.DataFrame(ruta_dias)

# =====================================================================
# CLIENTE GEMINI CON CORTACIRCUITOS RÁPIDO (SIN ESPERAS DE 30 MINUTOS)
# =====================================================================
MODELOS_RESPALDO = [
    "gemini-3-flash-preview",
    "gemini-3.0-flash",
    "gemini-3.1-flash-preview",
    "gemini-flash-latest",
    "gemini-2.5-flash",
    "gemini-3-pro-preview",
    "gemini-2.5-pro"
]

def obtener_cliente_gemini(api_key: str):
    if not api_key:
        return None
    return genai.Client(api_key=api_key.strip())

def obtener_cola_modelos(client, api_key: str):
    key_fp = hashlib.sha256(api_key.strip().encode("utf-8")).hexdigest()[:12]
    if key_fp in st.session_state.modelos_cache:
        return list(st.session_state.modelos_cache[key_fp])

    modelos_cuenta = []
    try:
        for m in client.models.list():
            nombre = getattr(m, "name", "")
            if not nombre:
                continue
            limpio = nombre.replace("models/", "")
            low = limpio.lower()
            if "gemini" not in low:
                continue
            excluidos = [
                "embedding", "imagen", "veo", "tts", "aqa",
                "audio", "live", "robotics", "computer-use", "learnlm"
            ]
            if any(x in low for x in excluidos):
                continue
            acciones = getattr(m, "supported_actions", None) or getattr(m, "supported_generation_methods", None)
            if acciones and not any("generatecontent" in str(a).lower() for a in acciones):
                continue
            modelos_cuenta.append(limpio)
    except Exception:
        pass

    def puntaje_modelo(mod_name):
        n = mod_name.lower()
        score = 0
        if "gemini-3" in n:
            score += 100
        elif "latest" in n:
            score += 80
        elif "gemini-2.5" in n:
            score += 50
        if "flash" in n and "lite" not in n:
            score += 25
        elif "lite" in n:
            score += 18
        elif "pro" in n:
            score += 15
        return -score

    ordenados = sorted(list(dict.fromkeys(modelos_cuenta)), key=puntaje_modelo)
    cola = ordenados + [m for m in MODELOS_RESPALDO if m not in ordenados]
    st.session_state.modelos_cache[key_fp] = cola
    return list(cola)

def llamar_gemini(client, api_key: str, prompt_contents, modelo_preferido="Auto-Detectar (Recomendado)", temperatura=0.1, max_modelos_probar=4):
    """Prueba únicamente los mejores modelos activos (máximo 4) con salto rápido."""
    cola_modelos = obtener_cola_modelos(client, api_key)
    if modelo_preferido and modelo_preferido != "Auto-Detectar (Recomendado)":
        cola_modelos = [modelo_preferido] + [m for m in cola_modelos if m != modelo_preferido]

    cola_corta = cola_modelos[:max_modelos_probar]
    errores = []

    for modelo in cola_corta:
        try:
            resp = client.models.generate_content(
                model=modelo,
                contents=prompt_contents,
                config=types.GenerateContentConfig(temperature=temperatura)
            )
            if resp and resp.text:
                return resp.text, modelo
        except Exception as e:
            msg = str(e).lower()
            errores.append(f"{modelo}: {str(e)[:120]}")
            if "429" in msg or "quota" in msg or "resource_exhausted" in msg:
                time.sleep(1.5)
                continue
            continue

    raise RuntimeError("No fue posible obtener respuesta de Gemini. Detalles: " + " | ".join(errores))

# =====================================================================
# MOTOR TURBO V4.0: EXTRACCIÓN EN SEGUNDOS (0 LLAMADAS API POR PÁGINA)
# =====================================================================
PATRON_BLOQUE_FECHA = re.compile(
    r'(?m)^(?P<fecha>\d{1,2}[\/\-]\d{1,2}[\/\-](?:\d{4}|\d{2}))[ \t]*(?:[-–:][ \t]*)?(?P<encabezado>[\w\t \(\)\.,#º°\-]{8,110})$'
)
PATRON_ACTO_FECHA = re.compile(
    r'(?i)(auto|providencia|resoluci[oó]n|sentencia|acta|oficio|notificaci[oó]n|audiencia|memorial|radicado|queja|recurso|orden|f[oó]rmula)[ \t]+(?:no\.?|n[o°]?)?[ \t]*[\w\-]*[ \t]*(?:del|de fecha|:)?[ \t]*(?P<fecha>\d{1,2}[\/\-]\d{1,2}[\/\-](?:\d{4}|\d{2}))'
)

def ocr_tesseract_turbo_gris(page) -> str:
    """
    OCR Ultraligero en Escala de Grises a 110 DPI (89.9% menos datos que 200 DPI RGB).
    - Descarta páginas en blanco en 0.001 segundos sin abrir Tesseract.
    - Usa LSTM directo (--oem 1 --psm 6) con timeout estricto de 4 segundos.
    - Gasta 0 llamadas de la cuota de Gemini API.
    """
    try:
        if not page.get_images() and not page.get_drawings():
            return ""

        pix = page.get_pixmap(dpi=110, colorspace=fitz.csGRAY, alpha=False)
        muestra = pix.samples[::20]
        if not muestra:
            del pix
            return ""

        pixeles_oscuros = sum(1 for b in muestra if b < 235)
        if (pixeles_oscuros / len(muestra)) < 0.003:
            del pix
            return ""

        img = Image.frombytes("L", [pix.width, pix.height], pix.samples)
        del pix

        import pytesseract
        texto_ocr = pytesseract.image_to_string(
            img,
            lang="spa",
            config="--oem 1 --psm 6",
            timeout=4
        ).strip()
        img.close()
        return texto_ocr if len(texto_ocr) > 15 else ""
    except Exception:
        return ""

def ejecutar_ocr_lote_gemini(client, api_key: str, lote_imagenes: list) -> dict:
    """Envía hasta 12 páginas escaneadas JUNTAS en 1 sola llamada a Gemini Flash."""
    if not lote_imagenes or client is None or not api_key:
        return {}

    contenidos = []
    instruccion = (
        "Transcribe el texto visible de cada una de las siguientes imágenes numeradas. "
        "Devuelve EXACTAMENTE cada bloque precedido por la etiqueta ===PAG_IDX_N=== (donde N es el número indicado). "
        "Sé literal con fechas, radicados, nombres y cifras. No agregues comentarios.\n"
    )
    contenidos.append(instruccion)
    for idx_item, (_, _, img_pil) in enumerate(lote_imagenes):
        contenidos.append(f"===PAG_IDX_{idx_item}===")
        contenidos.append(img_pil)

    try:
        texto_lote, _ = llamar_gemini(
            client=client,
            api_key=api_key,
            prompt_contents=contenidos,
            temperatura=0.0,
            max_modelos_probar=2
        )
        resultados = {}
        partes = re.split(r'===PAG_IDX_(\d+)===', texto_lote)
        for k in range(1, len(partes) - 1, 2):
            idx_recuperado = int(partes[k])
            resultados[idx_recuperado] = partes[k + 1].strip()
        return resultados
    except Exception:
        return {}

def limpiar_linea_meta(txt: str, max_len=110):
    return re.sub(r'\s+', ' ', str(txt)).strip()[:max_len]

def extraer_metadatos_documento(texto: str):
    muestra = texto[:50000]
    meta = {
        'radicado': 'No identificado',
        'actor': 'No identificado',
        'demandado': 'No identificado',
        'juez': 'No identificado',
        'tipo_actuacion': [],
        'numero_oficio': 'No identificado'
    }
    rr = RE_RADICADO.search(muestra)
    if rr:
        meta['radicado'] = limpiar_linea_meta(rr.group(0), 70)
    no_of = RE_NUMERO_OFICIO.search(muestra)
    if no_of:
        meta['numero_oficio'] = limpiar_linea_meta(no_of.group(0), 60)
    da = RE_ACTOR.search(muestra)
    if da:
        meta['actor'] = limpiar_linea_meta(da.group(2), 110)
    dd = RE_DEMANDADO.search(muestra)
    if dd:
        meta['demandado'] = limpiar_linea_meta(dd.group(2), 110)
    jj = RE_JUEZ.search(muestra)
    if jj:
        meta['juez'] = limpiar_linea_meta(jj.group(0), 110)

    texto_upper = muestra.upper()
    for tipo in TIPOS_ACT:
        if tipo in texto_upper:
            meta['tipo_actuacion'].append(tipo)
    return meta

def extraer_fechas_de_pagina(nombre_archivo: str, num_pag: int, txt_pag: str) -> list:
    hallazgos = []
    for m in PATRON_BLOQUE_FECHA.finditer(txt_pag):
        hallazgos.append({
            "Archivo": nombre_archivo,
            "Página": num_pag,
            "Fecha Literal": m.group("fecha").strip(),
            "Actuación / Referencia": re.sub(r'\s+', ' ', m.group("encabezado")).strip(),
            "Contexto": re.sub(r'\s+', ' ', txt_pag[max(0, m.start() - 25):min(len(txt_pag), m.end() + 120)]).strip()
        })
    for m in PATRON_ACTO_FECHA.finditer(txt_pag):
        hallazgos.append({
            "Archivo": nombre_archivo,
            "Página": num_pag,
            "Fecha Literal": m.group("fecha").strip(),
            "Actuación / Referencia": re.sub(r'\s+', ' ', m.group(0)).strip(),
            "Contexto": re.sub(r'\s+', ' ', txt_pag[max(0, m.start() - 25):min(len(txt_pag), m.end() + 120)]).strip()
        })
    return hallazgos

def procesar_expediente_turbo(archivos_subidos, api_key: str, client, modo_ocr: str, barra_progreso, texto_estado):
    t_inicio = time.time()
    docs_resultado = []
    cola_lote_nube = []
    presupuesto_ocr_local = 85
    ocr_local_usados = 0

    total_arch = len(archivos_subidos)

    for idx_arch, uploaded_file in enumerate(archivos_subidos):
        porcentaje = int(((idx_arch) / max(1, total_arch)) * 85)
        barra_progreso.progress(porcentaje)
        texto_estado.info(f"⚡ Lectura Turbo ({idx_arch + 1}/{total_arch}): **{uploaded_file.name}**...")

        file_bytes = uploaded_file.getvalue()
        hash_sha256 = hashlib.sha256(file_bytes).hexdigest()[:16]
        clave_cache = f"{hash_sha256}_{modo_ocr[:8]}"

        if clave_cache in st.session_state.cache_hashes:
            docs_resultado.append(st.session_state.cache_hashes[clave_cache])
            del file_bytes
            continue

        nombre = uploaded_file.name
        ext = Path(nombre).suffix.lower()
        bloques_paginas = []
        actuaciones_doc = []
        num_paginas = 0
        pags_ocr = 0

        try:
            if ext == ".pdf":
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                num_paginas = len(doc)
                for p_i in range(num_paginas):
                    page = doc[p_i]
                    num_pag = p_i + 1
                    txt = page.get_text("text").strip()

                    if len(txt) >= 35:
                        bloque = f"[Archivo: {nombre} - Pág. {num_pag}]\n{txt}"
                        bloques_paginas.append(bloque)
                        actuaciones_doc.extend(extraer_fechas_de_pagina(nombre, num_pag, bloque))
                    else:
                        if modo_ocr.startswith("🚀 Instantáneo"):
                            bloques_paginas.append(f"[Archivo: {nombre} - Pág. {num_pag} (Imagen/Anexo omitido en modo instantáneo)]")
                        elif modo_ocr.startswith("⚡ Turbo"):
                            if ocr_local_usados < presupuesto_ocr_local:
                                txt_ocr = ocr_tesseract_turbo_gris(page)
                                if txt_ocr:
                                    ocr_local_usados += 1
                                    pags_ocr += 1
                                    bloque = f"[Archivo: {nombre} - Pág. {num_pag} (OCR)]\n{txt_ocr}"
                                    bloques_paginas.append(bloque)
                                    actuaciones_doc.extend(extraer_fechas_de_pagina(nombre, num_pag, bloque))
                                else:
                                    bloques_paginas.append(f"[Archivo: {nombre} - Pág. {num_pag} (Sello/Página sin texto)]")
                            else:
                                bloques_paginas.append(f"[Archivo: {nombre} - Pág. {num_pag} (Omitida por tope preventivo CPU)]")
                        else:
                            if page.get_images() or page.get_drawings():
                                pix = page.get_pixmap(dpi=105, colorspace=fitz.csGRAY, alpha=False)
                                muestra = pix.samples[::20]
                                oscuros = sum(1 for b in muestra if b < 235) if muestra else 0
                                if muestra and (oscuros / len(muestra)) >= 0.003 and len(cola_lote_nube) < 48:
                                    img_pil = Image.frombytes("L", [pix.width, pix.height], pix.samples)
                                    cola_lote_nube.append((idx_arch, p_i, nombre, num_pag, img_pil))
                                    bloques_paginas.append("__PENDIENTE_OCR_LOTE__")
                                else:
                                    bloques_paginas.append(f"[Archivo: {nombre} - Pág. {num_pag} (Página en blanco)]")
                                del pix
                            else:
                                bloques_paginas.append(f"[Archivo: {nombre} - Pág. {num_pag} (Página en blanco)]")
                doc.close()

            elif ext == ".docx":
                doc = Document(io.BytesIO(file_bytes))
                parrafos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        celdas = [c.text.strip().replace("\n", " ") for c in row.cells]
                        if any(celdas):
                            parrafos.append(" - ".join(celdas))
                texto_docx = "\n".join(parrafos)
                tam_pag = 3000
                secciones = [texto_docx[k:k + tam_pag] for k in range(0, max(1, len(texto_docx)), tam_pag)]
                num_paginas = len(secciones)
                for idx_s, frag in enumerate(secciones, 1):
                    bloque = f"[Archivo: {nombre} - Sección {idx_s}]\n{frag}"
                    bloques_paginas.append(bloque)
                    actuaciones_doc.extend(extraer_fechas_de_pagina(nombre, idx_s, bloque))

            elif ext in (".txt", ".rtf"):
                contenido = file_bytes.decode("utf-8", errors="ignore").strip()
                num_paginas = 1
                bloque = f"[Archivo: {nombre} - Pág. 1]\n{contenido}"
                bloques_paginas.append(bloque)
                actuaciones_doc.extend(extraer_fechas_de_pagina(nombre, 1, bloque))

        except Exception as e:
            num_paginas = 1
            bloques_paginas.append(f"[Archivo: {nombre} - ERROR DE LECTURA: {str(e)}]")

        del file_bytes

        item_temp = {
            "nombre": nombre,
            "hash": hash_sha256,
            "clave_cache": clave_cache,
            "num_paginas": num_paginas,
            "pags_ocr": pags_ocr,
            "bloques_temp": bloques_paginas,
            "actuaciones": actuaciones_doc
        }
        docs_resultado.append(item_temp)

    if cola_lote_nube:
        tam_lote = 12
        total_lotes = math.ceil(len(cola_lote_nube) / tam_lote)
        for num_lote, ini in enumerate(range(0, len(cola_lote_nube), tam_lote), 1):
            sub_lote = cola_lote_nube[ini:ini + tam_lote]
            texto_estado.info(f"☁️ Procesando lote OCR en la nube {num_lote} de {total_lotes} ({len(sub_lote)} páginas juntas)...")
            mapa_ocr = ejecutar_ocr_lote_gemini(
                client=client,
                api_key=api_key,
                lote_imagenes=[(d_i, p_i, im) for (d_i, p_i, _, _, im) in sub_lote]
            )
            for idx_local, (d_i, p_i, nom_a, n_pag, im_pil) in enumerate(sub_lote):
                txt_rec = mapa_ocr.get(idx_local, "")
                im_pil.close()
                if txt_rec:
                    bloque_nuevo = f"[Archivo: {nom_a} - Pág. {n_pag} (OCR-Nube)]\n{txt_rec}"
                    docs_resultado[d_i]["bloques_temp"][p_i] = bloque_nuevo
                    docs_resultado[d_i]["pags_ocr"] += 1
                    docs_resultado[d_i]["actuaciones"].extend(extraer_fechas_de_pagina(nom_a, n_pag, bloque_nuevo))
                else:
                    docs_resultado[d_i]["bloques_temp"][p_i] = f"[Archivo: {nom_a} - Pág. {n_pag} (Sin texto legible)]"

    barra_progreso.progress(95)
    docs_finales = []
    for d in docs_resultado:
        if "bloques_temp" in d:
            texto_unificado = "\n\n".join(d["bloques_temp"])
            del d["bloques_temp"]
            clave_c = d.pop("clave_cache")
            d["caracteres"] = len(texto_unificado)
            d["meta"] = extraer_metadatos_documento(texto_unificado)
            d["raw"] = texto_unificado
            if len(st.session_state.cache_hashes) >= 40:
                st.session_state.cache_hashes.clear()
            st.session_state.cache_hashes[clave_c] = d
        docs_finales.append(d)

    barra_progreso.progress(100)
    duracion = round(time.time() - t_inicio, 1)
    return docs_finales, duracion

def consolidar_cronologia_df(items: list) -> pd.DataFrame:
    todas = []
    for it in items:
        todas.extend(it.get("actuaciones", []))
    if not todas:
        return pd.DataFrame(columns=["Fecha ISO", "Fecha Literal", "Archivo", "Página", "Actuación / Referencia", "Contexto"])

    def normalizar_fecha_rapida(f_str):
        s = str(f_str).strip().replace("-", "/")
        partes = s.split("/")
        if len(partes) == 3:
            try:
                dia, mes, anio = int(partes[0]), int(partes[1]), int(partes[2])
                if anio < 100:
                    anio += 2000 if anio < 50 else 1900
                if 1990 <= anio <= 2035 and 1 <= mes <= 12 and 1 <= dia <= 31:
                    return f"{anio:04d}-{mes:02d}-{dia:02d}"
            except Exception:
                pass
        return "9999-12-31"

    df_act = pd.DataFrame(todas)
    df_act["Fecha ISO"] = df_act["Fecha Literal"].apply(normalizar_fecha_rapida)
    df_act = df_act.drop_duplicates(subset=["Archivo", "Página", "Fecha Literal", "Actuación / Referencia"])
    df_act = df_act.sort_values(by=["Fecha ISO", "Archivo", "Página"]).reset_index(drop=True)
    df_act["Fecha ISO"] = df_act["Fecha ISO"].replace("9999-12-31", "Sin Fecha Exacta")
    return df_act[["Fecha ISO", "Fecha Literal", "Archivo", "Página", "Actuación / Referencia", "Contexto"]]

# =====================================================================
# GENERADOR DE WORD (.DOCX) Y EXCEL (.XLSX) BAJO DEMANDA
# =====================================================================
def aplicar_color_celda(cell, hex_color):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:val'), 'clear')
    shading_elm.set(qn('w:color'), 'auto')
    shading_elm.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def agregar_texto_con_negrillas(parrafo, texto):
    texto_limpio = texto.replace('`', '')
    partes = re.split(r'(\*\*.*?\*\*)', texto_limpio)
    for parte in partes:
        if parte.startswith('**') and parte.endswith('**') and len(parte) > 4:
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        else:
            parrafo.add_run(parte)

def dividir_fila_markdown_segura(linea_limpia: str):
    linea_sanitizada = re.sub(
        r'\[+[^\]]*\]+',
        lambda m: m.group(0).replace('|', ' - '),
        linea_limpia
    )
    partes = linea_sanitizada.strip().split('|')
    if linea_sanitizada.startswith('|'):
        partes = partes[1:]
    if linea_sanitizada.endswith('|'):
        partes = partes[:-1]
    return [c.strip() for c in partes]

def volcar_tabla_en_word(doc, filas_tabla):
    if not filas_tabla:
        return
    num_cols = len(filas_tabla[0])
    if num_cols == 0:
        return

    tabla = doc.add_table(rows=len(filas_tabla), cols=num_cols)
    tabla.style = 'Table Grid'

    for r_idx, fila in enumerate(filas_tabla):
        if len(fila) > num_cols:
            fila_ajustada = fila[:num_cols - 1] + [" - ".join(fila[num_cols - 1:])]
        else:
            fila_ajustada = fila + [""] * (num_cols - len(fila))

        for c_idx in range(num_cols):
            celda = tabla.cell(r_idx, c_idx)
            celda.text = ""
            p = celda.paragraphs[0]
            agregar_texto_con_negrillas(p, fila_ajustada[c_idx])

            if r_idx == 0:
                aplicar_color_celda(celda, "1E3A8A")
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)
            else:
                if r_idx % 2 == 1:
                    aplicar_color_celda(celda, "F8FAFC")
                for run in p.runs:
                    run.font.size = Pt(9)

    doc.add_paragraph()

def extraer_solo_borrador_radicable(md_texto: str) -> str:
    patron_inicio = re.compile(r'(?im)^##\s*(?:10|7)\.\s*BORRADOR.*$')
    m_ini = patron_inicio.search(md_texto)
    if not m_ini:
        return ""
    resto = md_texto[m_ini.end():].strip()
    patron_fin = re.compile(r'(?im)^##\s*(?:11|8)\.\s*CHECKLIST.*$')
    m_fin = patron_fin.search(resto)
    if m_fin:
        resto = resto[:m_fin.start()].strip()
    return resto

def convertir_markdown_a_docx_bytes(md_texto: str, titulo_portada="INFORME JURÍDICO Y PLAN DE ACCIÓN PROCESAL", incluir_portada=True) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    if incluir_portada:
        titulo = doc.add_heading(titulo_portada, level=0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = sub.add_run(f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')} | Bufete Litigante IA Colombia")
        run_sub.italic = True
        run_sub.font.size = Pt(9.5)
        run_sub.font.color.rgb = RGBColor(100, 116, 139)
        doc.add_paragraph("_" * 70)

    lineas = md_texto.split('\n')
    buffer_tabla = []

    for linea in lineas:
        linea_limpia = linea.strip()
        es_fila_tabla = linea_limpia.startswith('|') and linea_limpia.endswith('|') and len(linea_limpia) > 2

        if es_fila_tabla:
            contenido_sin_barras = linea_limpia.replace('|', '').replace('-', '').replace(':', '').strip()
            if not contenido_sin_barras:
                continue
            buffer_tabla.append(dividir_fila_markdown_segura(linea_limpia))
            continue
        else:
            if buffer_tabla:
                volcar_tabla_en_word(doc, buffer_tabla)
                buffer_tabla = []

        if not linea_limpia or linea_limpia == "---":
            continue

        if linea_limpia.startswith('# '):
            doc.add_heading(linea_limpia[2:].replace('**', ''), level=1)
        elif linea_limpia.startswith('## '):
            doc.add_heading(linea_limpia[3:].replace('**', ''), level=2)
        elif linea_limpia.startswith('### '):
            doc.add_heading(linea_limpia[4:].replace('**', ''), level=3)
        elif linea_limpia.startswith('#### '):
            doc.add_heading(linea_limpia[5:].replace('**', ''), level=4)
        elif linea_limpia.startswith(('- ', '* ')):
            p = doc.add_paragraph(style='List Bullet')
            agregar_texto_con_negrillas(p, linea_limpia[2:])
        elif re.match(r'^\d+\.\s+', linea_limpia):
            texto_item = re.sub(r'^\d+\.\s+', '', linea_limpia)
            p = doc.add_paragraph(style
